from __future__ import annotations

from dataclasses import replace
from datetime import datetime, timedelta, timezone
from pathlib import Path
from uuid import uuid4

from fastapi.testclient import TestClient

from adaptive_backend.app import create_app
from adaptive_backend.database import Database
from adaptive_backend.pedagogy import DynamicTeachingPolicy


def _client(settings, tmp_path: Path) -> TestClient:
    cfg = replace(settings, storage_dir=tmp_path / "uploads", auth_secret="s" * 48)
    return TestClient(create_app(cfg))


def _register(client: TestClient, username: str, role: str = "learner") -> tuple[str, dict[str, str]]:
    r = client.post(
        "/v1/auth/register",
        json={"username": username, "password": "correct horse battery staple", "role": role},
    )
    assert r.status_code == 200, r.text
    body = r.json()
    return body["user_id"], {"Authorization": f"Bearer {body['access_token']}"}


def _seed_misconception(db: Database, user_id: str, concept_id: str) -> None:
    now = datetime.now(timezone.utc)
    session_id = str(uuid4())
    with db.transaction() as conn:
        conn.execute(
            "INSERT INTO sessions(session_id,user_id,subject_id,started_at,status,source) VALUES(?,?,?,?,?,?)",
            (session_id, user_id, "BIO", (now - timedelta(days=2)).isoformat(), "completed", "system"),
        )
        for index, correct in enumerate((1, 0, 0, 1)):
            conn.execute(
                """
                INSERT INTO learning_events(
                    event_id,session_id,user_id,subject_id,primary_concept_id,concept_ids_json,
                    timestamp,activity_type,input_mode,correct,score,evidence_strength,mistake_type,metadata_json
                ) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?)
                """,
                (
                    str(uuid4()), session_id, user_id, "BIO", concept_id, db.dumps([concept_id]),
                    (now - timedelta(hours=8 - index)).isoformat(), "practice", "text", correct,
                    float(correct), 0.9, None if correct else "lysosome_peroxisome_confusion",
                    db.dumps({"hint_count": 0, "attempt_count": 1}),
                ),
            )
        conn.execute(
            """
            INSERT INTO concept_state(user_id,concept_id,mastery_belief,confidence,evidence_count,last_observed_at,status,recurrence_count,updated_at)
            VALUES(?,?,?,?,?,?,?,?,?)
            """,
            (user_id, concept_id, 0.43, 0.81, 4, now.isoformat(), "active_blocker", 2, now.isoformat()),
        )
        conn.execute(
            """
            INSERT INTO gap_hypotheses(user_id,target_concept_id,gap_concept_id,probability,source,evidence_json,updated_at)
            VALUES(?,?,?,?,?,?,?)
            """,
            (user_id, concept_id, "bio_lysosome", 0.76, "misconception", "[]", now.isoformat()),
        )


# ---------------------------------------------------------------------------
# Bug 1: causal cause reaches the teaching policy and is scoped to the target
# ---------------------------------------------------------------------------
def _programs(policy: dict) -> list[list[str]]:
    return [c["representation_sequence"] for c in [policy["selected"], *policy["alternatives"]]]


def test_policy_reads_causal_cause_type(core):
    db = core[0]
    policy = DynamicTeachingPolicy(db)
    for cause, first_step in (("misconception", "rival_models"), ("prerequisite", "prerequisite_bridge")):
        chosen = policy.choose(
            user_id="policy-user", subject_id="BIO", concept_id="bio_golgi",
            requested_action="explain", mode="study",
            diagnosis={"causal_cause_type": cause, "hypotheses": [{"source": "direct"}]},
        )
        assert chosen["context"]["diagnosis_source"] == cause
        assert any(seq and seq[0] == first_step for seq in _programs(chosen))


def test_kernel_cause_is_scoped_to_target(settings, tmp_path):
    with _client(settings, tmp_path) as client:
        user_id, _ = _register(client, "scopedlearner")
        services = client.app.state.services
        _seed_misconception(services.db, user_id, "bio_peroxisome")

        def plan(concept_id: str):
            return services.decision.plan(
                user_id=user_id, content_user_id=user_id, subject_id="BIO",
                concept_id=concept_id, requested_mode="practice", query=concept_id,
            )

        other = plan("bio_golgi")
        assert "causal_cause_type" not in other.diagnosis
        assert other.action != "contrast"
        # The subject-wide informational projection still reports the gap.
        assert other.diagnosis["causal"]["supported_case_count"] >= 1

        own = plan("bio_peroxisome")
        assert own.diagnosis.get("causal_cause_type") == "misconception"


# ---------------------------------------------------------------------------
# Bug 3: partially extracted uploads are fully set up; duplicates change nothing
# ---------------------------------------------------------------------------
def _two_slide_pptx() -> bytes:
    import io
    from pptx import Presentation

    prs = Presentation()
    first = prs.slides.add_slide(prs.slide_layouts[5])
    first.shapes.title.text = "The Golgi apparatus modifies, sorts and packages proteins."
    prs.slides.add_slide(prs.slide_layouts[6])  # blank slide -> no extractable text
    out = io.BytesIO()
    prs.save(out)
    return out.getvalue()


def test_partial_upload_gets_role_and_ir(settings, tmp_path):
    with _client(settings, tmp_path) as client:
        _, headers = _register(client, "partialowner")
        course = client.post("/v1/courses", headers=headers, json={"name": "Cells", "subject_name": "Cells"}).json()
        cid = course["course_id"]
        upload = client.post(
            f"/v1/documents/upload?course_id={cid}", headers=headers,
            files={"file": ("slides.pptx", _two_slide_pptx(),
                            "application/vnd.openxmlformats-officedocument.presentationml.presentation")},
        )
        assert upload.status_code == 200, upload.text
        record = upload.json()
        assert record["status"] == "ready_partial"
        db = client.app.state.services.db
        role = db.fetchone(
            "SELECT authority FROM document_source_roles WHERE course_id=? AND document_id=?",
            (cid, record["document_id"]),
        )
        assert role and role["authority"] == "canonical"
        ir = db.fetchone("SELECT count(*) AS n FROM learning_ir_objects WHERE document_id=?", (record["document_id"],))
        assert ir["n"] > 0


def test_duplicate_upload_keeps_teacher_role(settings, tmp_path):
    with _client(settings, tmp_path) as client:
        _, headers = _register(client, "dupowner")
        cid = client.post("/v1/courses", headers=headers, json={"name": "Dup", "subject_name": "Dup"}).json()["course_id"]
        body = b"Ribosomes assemble proteins from amino acids. Mitochondria produce ATP."
        files = {"file": ("notes.txt", body, "text/plain")}
        first = client.post(f"/v1/documents/upload?course_id={cid}", headers=headers, files=files).json()
        doc_id = first["document_id"]
        patched = client.patch(
            f"/v1/courses/{cid}/sources/{doc_id}", headers=headers,
            json={"authority": "supplementary", "explanation_weight": 0.4},
        )
        assert patched.status_code == 200, patched.text
        again = client.post(f"/v1/documents/upload?course_id={cid}", headers=headers, files=files).json()
        assert again["document_id"] == doc_id and again["duplicate"] is True
        role = client.app.state.services.db.fetchone(
            "SELECT authority,explanation_weight FROM document_source_roles WHERE course_id=? AND document_id=?",
            (cid, doc_id),
        )
        assert role["authority"] == "supplementary"
        assert abs(role["explanation_weight"] - 0.4) < 1e-9


# ---------------------------------------------------------------------------
# Bug 4: the learning map only exposes subjects the caller can access
# ---------------------------------------------------------------------------
def test_learning_map_hides_other_users_private_subjects(settings, tmp_path):
    with _client(settings, tmp_path) as client:
        _, owner = _register(client, "mapowner")
        _, stranger = _register(client, "mapstranger")
        course = client.post("/v1/courses", headers=owner, json={"name": "Secret Sci", "subject_name": "Secret Sci"}).json()
        concept = client.post(
            f"/v1/courses/{course['course_id']}/concepts", headers=owner, json={"name": "Hidden Topic Alpha"}
        ).json()
        private_subject = course["subject_id"]
        assert private_subject.startswith("USR:")

        def nodes(headers, **params):
            r = client.get("/v1/memory/map", headers=headers, params=params)
            assert r.status_code == 200, r.text
            return r.json()

        mine = nodes(owner)
        assert concept["concept_id"] in {n["id"] for n in mine["nodes"]}

        theirs = nodes(stranger)
        ids = {n["id"] for n in theirs["nodes"]}
        assert concept["concept_id"] not in ids
        assert "Hidden Topic Alpha" not in str(theirs)
        assert "bio_golgi" in ids  # built-in subjects stay visible

        guessed = nodes(stranger, subject_id=private_subject)
        assert guessed["nodes"] == [] and guessed["edges"] == []
        assert nodes(owner, subject_id=private_subject)["nodes"]


def test_learning_map_follows_class_enrollment(settings, tmp_path):
    with _client(settings, tmp_path) as client:
        _, teacher = _register(client, "mapteacher", "teacher")
        _, enrolled = _register(client, "mapenrolled")
        _, outsider = _register(client, "mapoutsider")
        course = client.post(
            "/v1/courses", headers=teacher,
            json={"name": "Class Bio", "subject_name": "Class Bio", "course_kind": "class"},
        ).json()
        cid = course["course_id"]
        concept = client.post(f"/v1/courses/{cid}/concepts", headers=teacher, json={"name": "Class Topic"}).json()
        assert client.post(f"/v1/courses/{cid}/enroll", headers=teacher, json={"learner_username": "mapenrolled"}).status_code == 200

        def ids(headers):
            return {n["id"] for n in client.get("/v1/memory/map", headers=headers).json()["nodes"]}

        assert concept["concept_id"] in ids(enrolled)
        assert concept["concept_id"] not in ids(outsider)
